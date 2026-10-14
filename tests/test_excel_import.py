import unittest
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from werkzeug.datastructures import FileStorage
from unittest.mock import patch

from app import app

import excel_export
import excel_import


class ExcelImportTests(unittest.TestCase):
    def _workbook_upload(self, image=False):
        workbook = Workbook()
        workstation = workbook.active
        workstation.title = "workstation"
        office = workbook.create_sheet("office")
        for sheet in (workstation, office):
            for column, (title, _key) in enumerate(excel_export.OFFICE_HEADERS, 1):
                sheet.cell(1, column, title)
        workstation.append([1, "YCSPGW000001", "HP", "Tower", "CPU", "Ethernet", "AA:BB", None, None, 512, 16, "SN-GW", None, None, None, None, "64-bit", "Windows", "24H2"])
        workstation.append([2, "YCSPGW000002"])
        workstation.append([3, "YCSPGW000003", "HP"])
        workstation.append([4, None, "HP", None, None, None, None, None, None, None, None, "SN-NO-MACHINE"])
        office.append([1, "YCSPBG000001", "Lenovo", "Laptop", "CPU", "Ethernet", "CC:DD", "Wi-Fi", "EE:FF", 512, 16, "SN-BG", "IT", "Alice", None, "2026-01-01", "64-bit", "Windows", "24H2"])
        if image:
            image_bytes = BytesIO()
            Image.new("RGB", (3, 3), color="red").save(image_bytes, format="PNG")
            image_bytes.seek(0)
            workstation.add_image(ExcelImage(image_bytes), "T2")
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        return FileStorage(stream=stream, filename="inventory.xlsx")

    def test_reads_workstation_and_office_sheets_without_images(self):
        entries = excel_import.read_excel_records(self._workbook_upload())
        self.assertEqual(len(entries), 2)
        self.assertEqual(entries[0]["fields"]["machine_name"], "YCSPGW000001")
        self.assertEqual(entries[1]["fields"]["department"], "IT")
        self.assertIsNone(entries[0]["image"])

    def test_reads_optional_embedded_image(self):
        entries = excel_import.read_excel_records(self._workbook_upload(image=True))
        self.assertEqual(entries[0]["image"][0], "png")
        self.assertTrue(entries[0]["image"][1].startswith(b"\x89PNG"))
        self.assertIsNone(entries[1]["image"])

    def test_import_route_persists_rows_as_unreviewed_records(self):
        client = app.test_client()
        with client.session_transaction() as session:
            session['user'] = 'tester'
        upload = self._workbook_upload()
        with patch('app.database.find_record_by_machine_name', return_value=None), \
             patch('app.database.find_record_by_sn', return_value=None), \
             patch('app.database.save_record', side_effect=[101, 102]) as save_record:
            response = client.post(
                '/import/excel',
                data={'excel_file': (upload.stream, upload.filename)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(save_record.call_count, 2)
        imported_fields = [call.kwargs['parsed'] for call in save_record.call_args_list]
        self.assertEqual([item['machine_name'] for item in imported_fields], [
            'YCSPGW000001', 'YCSPBG000001',
        ])
        self.assertEqual([item['migrate_domain'] for item in imported_fields], ['是', '是'])
        self.assertTrue(all(call.kwargs['image_path'] == '' for call in save_record.call_args_list))

    def test_import_route_stores_embedded_images(self):
        from tempfile import TemporaryDirectory
        from pathlib import Path

        client = app.test_client()
        with client.session_transaction() as session:
            session['user'] = 'tester'
        upload = self._workbook_upload(image=True)
        with TemporaryDirectory() as upload_dir, \
             patch('app.Config.UPLOAD_FOLDER', upload_dir), \
             patch('app.database.find_record_by_machine_name', return_value=None), \
             patch('app.database.find_record_by_sn', return_value=None), \
             patch('app.database.save_record', side_effect=[101, 102]) as save_record:
            response = client.post(
                '/import/excel',
                data={'excel_file': (upload.stream, upload.filename)},
                content_type='multipart/form-data',
                follow_redirects=True,
            )
            self.assertEqual(response.status_code, 200)
            first_image_path = save_record.call_args_list[0].kwargs['image_path']
            self.assertTrue(Path(first_image_path).is_file())
            self.assertTrue(Path(first_image_path).read_bytes().startswith(b'\x89PNG'))
            self.assertEqual(save_record.call_args_list[1].kwargs['image_path'], '')
    def test_rejects_non_xlsx_files(self):
        upload = FileStorage(stream=BytesIO(b"not excel"), filename="inventory.xls")
        with self.assertRaisesRegex(ValueError, "xlsx"):
            excel_import.read_excel_records(upload)

    def test_rejects_workbook_without_machine_name_header(self):
        workbook = Workbook()
        workbook.active.append(["unrelated", "columns"])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        upload = FileStorage(stream=stream, filename="wrong.xlsx")
        with self.assertRaisesRegex(ValueError, "计算机名"):
            excel_import.read_excel_records(upload)


if __name__ == "__main__":
    unittest.main()
